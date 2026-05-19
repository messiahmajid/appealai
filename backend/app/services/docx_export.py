import io
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _clean_for_submission(letter: str) -> str:
    text = letter
    text = re.sub(r'\s*\[SOURCE [A-Z]\]', '', text)
    text = re.sub(
        r'^(\s{0,3}(?:#{1,6}\s*)?(?:\*\*)?)Documentation Gaps(?:\*\*)?\s*$',
        r'\1Clinical Rationale for Exception',
        text,
        flags=re.IGNORECASE | re.MULTILINE,
    )
    text = re.sub(
        r'\[DOCUMENTATION GAP[.:]\s*(.*?)\]',
        r'Additional supporting rationale: \1',
        text,
        flags=re.DOTALL,
    )
    text = re.sub(r'\[DOCUMENTATION GAP\]\s*:\s*', 'Additional supporting rationale: ', text)
    text = re.sub(r'\[DOCUMENTATION GAP[.:]\s*', 'Additional supporting rationale: ', text)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^---+\s*$', '', text, flags=re.MULTILINE)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def generate_docx(letter: str, patient_name: str = "") -> io.BytesIO:
    clean = _clean_for_submission(letter)
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.15

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    for line in clean.split('\n'):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph('')
            continue

        is_header = (
            stripped.startswith('RE:')
            or stripped.startswith('Subject:')
            or stripped == 'References'
            or stripped == 'Guidelines'
            or stripped == 'Literature'
            or stripped in (
                'Clinical Summary',
                'Criterion-by-Criterion Medical Necessity Argument',
                'Direct Rebuttal to Denial Reason',
                'Direct Denial Rebuttal',
                'Clinical Rationale for Exception',
                'Clinical Rationale for Step-Therapy Exception',
                'Additional Supporting Rationale',
                'Supporting Documentation',
                'Conclusion',
            )
        )

        p = doc.add_paragraph()

        if is_header:
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p.style = doc.styles['List Bullet']
            p.add_run(stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            p.add_run(stripped)
            p.paragraph_format.left_indent = Inches(0.25)
        else:
            p.add_run(stripped)
            if any(stripped.startswith(k) for k in ('Patient Name:', 'Date of Birth:', 'Member ID:', 'CPT Code', 'ICD-10', 'Claim/', 'Date of', 'Denied Service:', 'Physician:')):
                p.runs[0].bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
